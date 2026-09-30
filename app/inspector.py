import io
import zipfile

import docx
import pymupdf
from PIL import Image, UnidentifiedImageError

from app.markers.base import FilePayload

MIN_TYPED_CHARS = 50  # a PDF with less text than this counts as scanned
DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
IMAGE_MIMES = {"PNG": "image/png", "JPEG": "image/jpeg", "WEBP": "image/webp"}


# Looks at a file and decides: typed or handwritten.
# Typed files come back with their text filled in, handwritten with None.
def inspect_file(data: bytes) -> FilePayload:
    if data[:4] == b"%PDF":
        return _inspect_pdf(data)
    if zipfile.is_zipfile(io.BytesIO(data)):
        return _inspect_docx(data)
    return _inspect_image(data)


# PDF with real text inside = typed. No text = scanned/handwritten.
def _inspect_pdf(data: bytes) -> FilePayload:
    doc = pymupdf.open(stream=data, filetype="pdf")
    text = "\n".join(page.get_text() for page in doc).strip()
    if len(text) >= MIN_TYPED_CHARS:
        return FilePayload(data, "application/pdf", text)
    return FilePayload(data, "application/pdf", None)


# Word file = typed. Reads paragraphs and table cells.
def _inspect_docx(data: bytes) -> FilePayload:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if "word/document.xml" not in z.namelist():
            raise ValueError("Unsupported file type")
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return FilePayload(data, DOCX_MIME, "\n".join(parts).strip())


# Photo = handwritten. Only PNG, JPEG and WEBP are accepted.
def _inspect_image(data: bytes) -> FilePayload:
    try:
        fmt = Image.open(io.BytesIO(data)).format
    except (UnidentifiedImageError, OSError):
        raise ValueError("Unsupported file type")
    if fmt not in IMAGE_MIMES:
        raise ValueError("Unsupported file type")
    return FilePayload(data, IMAGE_MIMES[fmt], None)
